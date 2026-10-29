# -*- coding: utf-8 -*-
"""
Genera la nomina de docentes a partir de los cuadros de personal (PDF).

- El usuario selecciona uno o VARIOS PDFs del cuadro de personal y el lugar
  donde guardar el Excel (ventanas graficas; no requiere escribir rutas).
- Extrae TODOS los docentes y catedraticos, SIN DUPLICADOS (una fila por
  persona, identificada por cedula, aunque aparezca en varios PDFs/turnos).
- Detalla TURNO y CATEGORIA de sueldo (rubro, ej. ZZ2).
- Clasificacion por categoria de rubro:
    L*  = docente   ZZ* = docente catedratico   G* = funcionario no docente
    (los G* se excluyen; para incluirlos usar --todos o INCLUIR_FUNCIONARIOS)
- El nombre y apellido se formatea como en la nomina:
    APELLIDO(S), NOMBRE(S)     ej.  PEREZ GOMEZ, JUAN CARLOS
  La separacion es automatica; solo los casos realmente ambiguos (nombres
  compuestos o apellidos compuestos) van en la tabla MANUAL (clave = cedula).

Modos de uso:
  - Sin argumentos: modo grafico (elegir PDFs y donde guardar el Excel).
  - "generar_nomina.py --cli": usa los PDFs de CUADRO_PERSONAL_PDF/ y guarda
    el resultado junto al script.
  - "generar_nomina.py --pdfs a.pdf b.pdf --out salida.xlsx": modo consola
    con archivos y destino indicados.
"""
import os
import queue
import re
import subprocess
import sys
import threading
import time
import traceback
import unicodedata
from datetime import datetime

try:
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk
except ImportError:   # entornos sin Tk: el modo consola debe seguir funcionando
    tk = None
    filedialog = messagebox = ttk = None

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

BASE = os.path.dirname(os.path.abspath(__file__))

def _ruta_icono():
    """Ruta del icono de la aplicacion (funciona en script y en EXE)."""
    base = getattr(sys, "_MEIPASS", BASE)   # dentro del EXE onefile
    return os.path.join(base, "favicon.ico")

# ----------------------------------------------------------------------
# 1. Configuracion
# ----------------------------------------------------------------------
VERSION = "1.1.0"
REPO = "rodrigoangeloni/nomina_docente_de_cuadro_perdonal"   # para chequear actualizaciones
INCLUIR_FUNCIONARIOS = False   # True -> incluye tambien funcionarios (G*)
ORDEN_TURNO = ["M", "T", "N", "TN"]
CED_RE = re.compile(r"^[0-9]{4,8}$")

# Ajustes manuales de nombres ambiguos (opcional). Si existe el archivo local
# manual_nombres.py (excluido del repo por contener datos personales) se usa;
# si no, todo se resuelve con la separacion automatica.
try:
    from manual_nombres import MANUAL
except ImportError:
    MANUAL = {}

CONNECTORES = {"DE", "DEL"}

# Identidad visual y textos de la interfaz grafica
NOMBRE_APP = "Nómina de Docentes"
AUTOR = "Rodrigo Angeloni"
COLOR_AZUL = "#2E75B6"
COLOR_AZUL_OSCURO = "#1F4E79"
COLOR_FONDO = "#F4F7FB"
COLOR_BORDE = "#D6E0EA"
COLOR_TEXTO = "#1F2933"
COLOR_TEXTO_SUAVE = "#5B6B7C"
COLOR_VERDE = "#1E7B45"
COLOR_VERDE_CLARO = "#E8F5EE"
COLOR_GRIS = "#8A97A6"

# ----------------------------------------------------------------------
# 2. Formato de nombres
# ----------------------------------------------------------------------
def separar_apellidos_nombre(bruto):
    """'MARIA ELENA GARCIA LOPEZ' -> 'GARCIA LOPEZ, MARIA ELENA' (heuristica).
    Los apellidos paterno y materno van al final; los conectores DE / DEL
    pertenecen al bloque de apellidos."""
    t = bruto.upper().split()
    n = len(t)
    if n <= 1:
        return t[0] if t else ""
    # bloque de apellidos: por defecto los ultimos 2 token (1 si solo hay 2)
    start = n - 2 if n >= 3 else n - 1
    cambio = True
    while cambio and start > 0:
        cambio = False
        if t[start] in CONNECTORES and start - 1 > 0:
            # el bloque empieza con conector -> el apellido previo tambien es apellido
            start -= 1
            cambio = True
    return " ".join(t[start:]) + ", " + " ".join(t[:start])

def formatear_nombre(bruto, cedula):
    if cedula in MANUAL:
        return MANUAL[cedula]
    return separar_apellidos_nombre(bruto)

# ----------------------------------------------------------------------
# 3. Extraccion de filas e institucion del cuadro
# ----------------------------------------------------------------------
def _normalizar(texto):
    """minusculas sin acentos, para comparar etiquetas del PDF."""
    texto = unicodedata.normalize("NFKD", texto or "")
    return "".join(c for c in texto if not unicodedata.combining(c)).lower().strip()

def _lineas_palabras(palabras, tolerancia=3.0):
    """Agrupa palabras de pdfplumber por linea visual."""
    lineas = []
    for palabra in sorted(palabras, key=lambda w: (w["top"], w["x0"])):
        if not lineas or abs(palabra["top"] - lineas[-1]["top"]) > tolerancia:
            lineas.append({"top": palabra["top"], "palabras": [palabra]})
        else:
            lineas[-1]["palabras"].append(palabra)
    for linea in lineas:
        linea["palabras"].sort(key=lambda w: w["x0"])
    return lineas

def extraer_institucion(pdf_path):
    """Lee el encabezado del cuadro y devuelve nombre, codigo, telefono y direccion.

    El formulario del PDF tiene los campos 3 (nombre), 4 (codigo), 5 (telefono)
    y 6 (direccion) en la parte superior de la primera pagina. Se ubican por las
    coordenadas de las etiquetas para no depender del orden del texto extraido.
    """
    with pdfplumber.open(pdf_path) as pdf:
        if not pdf.pages:
            return {}
        pagina = pdf.pages[0]
        ancho = pagina.width
        palabras = pagina.extract_words()

    lineas = _lineas_palabras(palabras)
    indice = None
    for i, linea in enumerate(lineas):
        texto = _normalizar(" ".join(w["text"] for w in linea["palabras"]))
        if "nombre de la institucion" in texto:
            indice = i
            break
    if indice is None:
        return {}

    linea_label = lineas[indice]
    # La banda de valores termina en la siguiente linea de etiquetas numeradas
    # (por ejemplo "7. Localidad:"), no en la primera linea de valores.
    fin = linea_label["top"] + 35
    for linea in lineas[indice + 1:]:
        texto = " ".join(w["text"] for w in linea["palabras"])
        if re.match(r"^\s*\d+\.", texto) and linea["top"] > linea_label["top"] + 10:
            fin = linea["top"]
            break
    x_codigo = x_telefono = x_direccion = None
    for palabra in linea_label["palabras"]:
        etiqueta = palabra["text"].rstrip(".")
        if etiqueta == "4":
            x_codigo = palabra["x0"]
        elif etiqueta == "5":
            x_telefono = palabra["x0"]
        elif etiqueta == "6":
            x_direccion = palabra["x0"]
    if x_codigo is None:
        x_codigo = ancho * 0.37
    if x_telefono is None:
        x_telefono = ancho * 0.53
    if x_direccion is None:
        x_direccion = ancho * 0.65

    banda = [w for w in palabras if linea_label["top"] < w["top"] < fin]

    def campo(desde, hasta):
        elegidas = [w for w in banda if desde <= (w["x0"] + w["x1"]) / 2 < hasta]
        elegidas.sort(key=lambda w: (round(w["top"], 1), w["x0"]))
        return re.sub(r"\s+", " ", " ".join(w["text"] for w in elegidas)).strip()

    # El campo de direccion a veces arrastra un corchete de la localidad
    # (por ejemplo "[MBACHIO]"); se recorta para no ensuciar la cabecera.
    direccion = re.sub(r"\s*\[.*$", "", campo(x_direccion - 4, ancho + 1)).strip()

    return {
        "nombre": campo(0, x_codigo - 4),
        "codigo": campo(x_codigo - 4, x_telefono - 4),
        "telefono": campo(x_telefono - 4, x_direccion - 4),
        "direccion": direccion,
    }

def extraer_instituciones(pdf_paths):
    """Devuelve las instituciones unicas detectadas en los PDFs seleccionados."""
    instituciones = []
    vistos = set()
    for path in pdf_paths:
        try:
            info = extraer_institucion(path)
        except Exception:
            continue
        nombre = (info.get("nombre") or "").strip()
        if not nombre:
            continue
        clave = _normalizar(nombre)
        if clave in vistos:
            continue
        vistos.add(clave)
        info["archivo"] = os.path.basename(path)
        instituciones.append(info)
    return instituciones

def extract_filas(pdf_path):
    """Devuelve (nombre_bruto, cedula, categoria_rubro, turno) por cada fila
    de persona de las tablas del PDF."""
    filas = []
    with pdfplumber.open(pdf_path) as pdf:
        for p in pdf.pages:
            for t in p.extract_tables():
                for r in t:
                    cells = [str(c).replace("\n", " ").strip() if c is not None else "" for c in r]
                    if len(cells) < 12:
                        continue
                    nombre, ced, cat = cells[1], cells[2], cells[3]
                    turno = cells[10]
                    if not nombre or not ced:
                        continue
                    if nombre.upper().startswith("NOMBRES") or re.match(r"^[0-9.\s-]+$", nombre):
                        continue
                    if not CED_RE.match(ced):
                        continue
                    filas.append((nombre, ced, cat, turno))
    return filas

# ----------------------------------------------------------------------
# 4. Procesamiento principal
# ----------------------------------------------------------------------
def tipo(cats):
    for c in cats:
        if c.startswith("Z"):   # ZZ* = catedratico
            return "Catedratico"
        if c.startswith("L"):   # L* = docente
            return "Docente"
        if c.startswith("G"):   # G* = funcionario no docente
            return "Funcionario"
    return "Otro"

def procesar(pdf_paths, out_path, incluir_funcionarios=INCLUIR_FUNCIONARIOS, progreso=None):
    """Extrae, consolida y genera el Excel. Devuelve un resumen (dict).

    `progreso`, si se indica, se llama como progreso(actual, total, mensaje)
    antes de leer cada PDF y antes de escribir el Excel. La usa la interfaz
    grafica para mostrar el avance; la consola puede omitirla.
    """
    # dedupe por cedula normalizada (sin ceros a la izquierda)
    personas = {}
    total_pdfs = len(pdf_paths)
    if progreso:
        try:
            progreso(0, total_pdfs, "Leyendo los datos de las instituciones...")
        except Exception:
            pass
    instituciones = extraer_instituciones(pdf_paths)
    for indice, path in enumerate(pdf_paths, 1):
        if progreso:
            try:
                progreso(indice - 1, total_pdfs, "Leyendo %s..." % os.path.basename(path))
            except Exception:
                pass   # un fallo al informar no debe cortar la generacion
        for nombre, ced, cat, turno in extract_filas(path):
            clave = str(int(ced))
            p = personas.setdefault(clave, {"nombre": nombre, "turnos": [], "cats": []})
            if cat and cat.upper() not in p["cats"]:
                p["cats"].append(cat.upper())
            if turno and turno.upper() not in p["turnos"]:
                p["turnos"].append(turno.upper())

    # si ademas de 'Docente Director' hay categoria de rubro, quedan solo rubros
    for p in personas.values():
        rubros = [c for c in p["cats"] if re.match(r"^[LGZ]", c)]
        if rubros:
            p["cats"] = rubros
        p["turnos"].sort(key=lambda t: ORDEN_TURNO.index(t) if t in ORDEN_TURNO else len(ORDEN_TURNO))

    lista = []
    excluidos = 0
    for ced, p in personas.items():
        t = tipo(p["cats"])
        if t in ("Funcionario", "Otro") and not incluir_funcionarios:
            excluidos += 1
            continue
        lista.append({
            "nombre": formatear_nombre(p["nombre"], ced),
            "ced": ced,
            "turno": " / ".join(p["turnos"]),
            "categoria": " / ".join(p["cats"]),
        })
    lista.sort(key=lambda x: x["nombre"].upper())

    tot = {"Docente": 0, "Catedratico": 0, "Funcionario": 0, "Otro": 0}
    for x in lista:
        tot[tipo(x["categoria"].split(" / "))] += 1

    if progreso:
        try:
            progreso(total_pdfs, total_pdfs, "Escribiendo el Excel...")
        except Exception:
            pass
    crear_excel(lista, out_path, instituciones)

    return {
        "total": len(lista),
        "instituciones": instituciones,
        "docentes": tot["Docente"],
        "catedraticos": tot["Catedratico"],
        "funcionarios": tot["Funcionario"],
        "excluidos": excluidos,
        "otros": [x for x in lista if tipo(x["categoria"].split(" / ")) == "Otro"],
        "destino": out_path,
    }

# ----------------------------------------------------------------------
# 5. Excel con el estilo de la nomina oficial
# ----------------------------------------------------------------------
def crear_excel(lista, out_path, instituciones=None):
    hoy = datetime.now()

    nombres = []
    direcciones = []
    vistos_nombres = set()
    vistos_direcciones = set()
    for info in instituciones or []:
        nombre = (info.get("nombre") or "").strip()
        if nombre:
            clave = _normalizar(nombre)
            if clave not in vistos_nombres:
                vistos_nombres.add(clave)
                nombres.append(nombre)
        direccion = (info.get("direccion") or "").strip()
        if direccion:
            clave = _normalizar(direccion)
            if clave not in vistos_direcciones:
                vistos_direcciones.add(clave)
                direcciones.append(direccion)
    if not nombres:
        nombres = ["NÓMINA DE DOCENTES"]
    texto_nombres = "\n".join(nombres)
    texto_direcciones = "\n".join(direcciones)

    lineas_nombres = sum(max(1, (len(nombre) + 79) // 80) for nombre in nombres)
    lineas_direcciones = sum(
        max(1, (len(direccion) + 84) // 85) for direccion in direcciones) or 1

    wb = Workbook()
    ws = wb.active
    ws.title = "NOMINA DOCENTE " + str(hoy.year)

    ws.column_dimensions["A"].width = 6.14
    ws.column_dimensions["B"].width = 56.86
    ws.column_dimensions["C"].width = 22.57
    ws.column_dimensions["D"].width = 9.0
    ws.column_dimensions["E"].width = 13.5

    ws.row_dimensions[1].height = 36.0
    ws.row_dimensions[2].height = max(40.5, lineas_nombres * 17 + 6)
    ws.row_dimensions[3].height = max(15.0, lineas_direcciones * 13 + 2)
    ws.row_dimensions[4].height = 5.25
    ws.row_dimensions[5].height = 20.25
    ws.row_dimensions[6].height = 7.5
    ws.row_dimensions[7].height = 39.75

    ws.merge_cells("A2:C2")
    ws["A2"] = texto_nombres
    ws["A2"].font = Font(name="Times New Roman", size=12, bold=True)
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A3:C3")
    ws["A3"] = texto_direcciones
    ws["A3"].font = Font(name="Times New Roman", size=9, bold=True)
    ws["A3"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.merge_cells("A5:C5")
    ws["A5"] = "NOMINA DE DOCENTES %02d/%d" % (hoy.month, hoy.year)
    ws["A5"].font = Font(name="Times New Roman", size=12, bold=True)
    ws["A5"].alignment = Alignment(horizontal="center", vertical="center")

    headers = ["N°", "APELLIDO Y NOMBRE", "CEDULA DE\nIDENTIDAD", "TURNO", "CATEGORIA\nDE SUELDO"]
    header_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2E75B6")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin = Side(style="thin", color="000000")
    head_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for col, val in zip(("A", "B", "C", "D", "E"), headers):
        c = ws[col + "7"]
        c.value = val
        c.font = header_font
        c.fill = header_fill
        c.alignment = header_align
        c.border = head_border

    def ced_val(ced):
        return int(ced) if ced.isdigit() else ced

    start = 8
    data_border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, x in enumerate(lista):
        r = start + i
        ws.cell(r, 1, value=i + 1)
        ws.cell(r, 2, value=x["nombre"])
        ws.cell(r, 3, value=ced_val(x["ced"]))
        ws.cell(r, 4, value=x["turno"])
        ws.cell(r, 5, value=x["categoria"])
        ws.row_dimensions[r].height = 20.1
        for col, horiz in ((1, "center"), (2, "left"), (3, "center"), (4, "center"), (5, "center")):
            c = ws.cell(r, col)
            c.border = data_border
            c.alignment = Alignment(horizontal=horiz, vertical="center")
            c.font = Font(name="Calibri", size=11 if col == 1 else 12)

    wb.save(out_path)

# ----------------------------------------------------------------------
# 6. Chequeo de actualizaciones (solo cuando corre como EXE)
# ----------------------------------------------------------------------
# GitHub limita su API (60 consultas/hora por IP sin autenticar), por eso
# el chequeo se hace A LO SUMO una vez cada INTERVALO_CHEQUEO_HORAS y se
# guarda en cache local (tambien la version que el usuario decidio omitir).
INTERVALO_CHEQUEO_HORAS = 24

def _version_tuple(v):
    nums = re.findall(r"\d+", v)[:3]
    return tuple(int(x) for x in nums) if nums else (0,)

def _ruta_cache():
    base = os.environ.get("LOCALAPPDATA") or os.path.expanduser("~")
    carpeta = os.path.join(base, "NominaDocentes")
    os.makedirs(carpeta, exist_ok=True)
    return os.path.join(carpeta, "actualizacion.json")

def _leer_cache():
    try:
        import json
        with open(_ruta_cache(), encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}

def _guardar_cache(data):
    try:
        import json
        with open(_ruta_cache(), "w", encoding="utf-8") as f:
            json.dump(data, f)
    except Exception:
        pass   # la cache nunca debe romper el programa

def _consultar_github():
    """(tag, url_exe) del ultimo release publicado."""
    import json
    import urllib.request
    req = urllib.request.Request(
        "https://api.github.com/repos/%s/releases/latest" % REPO,
        headers={"User-Agent": "nomina-docentes"},
    )
    with urllib.request.urlopen(req, timeout=5) as r:
        data = json.load(r)
    tag = (data.get("tag_name") or "").lstrip("vV")
    asset = next((a for a in data.get("assets", []) if a.get("name", "").lower().endswith(".exe")), None)
    return tag, (asset["browser_download_url"] if asset else None)

def hay_nueva_version():
    """Devuelve (tag, url_exe) si hay version mas nueva que VERSION pendiente.
    Consulta la red solo si paso el intervalo; si no, usa la cache."""
    if not getattr(sys, "frozen", False):   # solo tiene sentido en el EXE
        return None

    cache = _leer_cache()
    ahora = time.time()
    vence = cache.get("last_check", 0) + INTERVALO_CHEQUEO_HORAS * 3600

    if ahora < vence:
        # dentro del intervalo: usar lo ya consultado, sin tocar la red
        tag, url = cache.get("last_tag"), cache.get("last_url")
        if tag and url and _version_tuple(tag) > _version_tuple(VERSION) and tag != cache.get("declined"):
            return tag, url
        return None

    # paso el intervalo: consultar la red (una sola vez cada X horas)
    try:
        tag, url = _consultar_github()
    except Exception:
        cache["last_check"] = ahora   # sin red: reintentar recien pasado el intervalo
        _guardar_cache(cache)
        return None

    cache["last_check"] = ahora
    cache["last_tag"], cache["last_url"] = tag, url
    _guardar_cache(cache)

    if tag and url and _version_tuple(tag) > _version_tuple(VERSION) and tag != cache.get("declined"):
        return tag, url
    return None

def descargar_y_reemplazar(url):
    """Descarga el EXE nuevo y reemplaza el actual (se puede renombrar un exe
    en ejecucion, no sobrescribir). Devuelve la ruta del .old."""
    import shutil
    import urllib.request
    exe = sys.executable
    nuevo = exe + ".nueva"
    req = urllib.request.Request(url, headers={"User-Agent": "nomina-docentes"})
    with urllib.request.urlopen(req, timeout=120) as r, open(nuevo, "wb") as f:
        shutil.copyfileobj(r, f)
    viejo = exe + ".old"
    if os.path.exists(viejo):
        os.remove(viejo)
    os.rename(exe, viejo)
    shutil.move(nuevo, exe)
    # limpiar la cache: el exe nuevo arranca con datos frescos
    try:
        os.remove(_ruta_cache())
    except Exception:
        pass
    return viejo

def verificar_actualizacion(root):
    """Pregunta al usuario si quiere actualizar. Si acepta, descarga, reemplaza
    y reinicia. Devuelve True si se actualizo (y hay que salir)."""
    try:
        upd = hay_nueva_version()
    except Exception:
        return False          # sin internet o sin release: seguir normal
    if not upd:
        return False
    tag, url = upd
    si = messagebox.askyesno(
        "Nueva versión disponible",
        "Hay una versión nueva: v%s\nTu versión actual: v%s\n\n"
        "¿Querés descargarla ahora?\n\n(Si elegís No, no se volverá a preguntar por esta versión)" % (tag, VERSION),
        parent=root,
    )
    if not si:
        cache = _leer_cache()
        cache["declined"] = tag   # no volver a preguntar por esta version
        _guardar_cache(cache)
        return False
    try:
        descargar_y_reemplazar(url)
    except Exception as e:
        messagebox.showerror(
            "Actualización",
            "No se pudo completar la actualización:\n\n%s\n\nSe continuará con la versión actual." % e,
            parent=root,
        )
        return False
    messagebox.showinfo("Actualización", "Actualizado a v%s. El programa se reiniciará." % tag, parent=root)
    subprocess.Popen([sys.executable])
    return True

# ----------------------------------------------------------------------
# 7. Modo grafico (interfaz guiada para usuarios no tecnicos)
# ----------------------------------------------------------------------
def _carpeta_sugerida():
    """Carpeta Documentos del usuario (o el home si no existe)."""
    home = os.path.expanduser("~")
    for nombre in ("Documents", "Documentos",
                   os.path.join("OneDrive", "Documents"),
                   os.path.join("OneDrive", "Documentos")):
        ruta = os.path.join(home, nombre)
        if os.path.isdir(ruta):
            return ruta
    return home

def _nombre_sugerido():
    hoy = datetime.now()
    return "NOMINA DOCENTES %02d-%d.xlsx" % (hoy.month, hoy.year)

def _centrar(ventana, ancho, alto, padre=None):
    ventana.update_idletasks()
    if padre is not None and padre.winfo_exists():
        x = padre.winfo_rootx() + max(0, (padre.winfo_width() - ancho) // 2)
        y = padre.winfo_rooty() + max(0, (padre.winfo_height() - alto) // 2)
    else:
        x = max(0, (ventana.winfo_screenwidth() - ancho) // 2)
        y = max(0, (ventana.winfo_screenheight() - alto) // 2 - 20)
    ventana.geometry("%dx%d+%d+%d" % (ancho, alto, x, y))

def _abrir_con_sistema(ruta, carpeta=False):
    """Abre un archivo o su carpeta con la aplicacion predeterminada."""
    if sys.platform.startswith("win"):
        os.startfile(os.path.dirname(ruta) if carpeta else ruta)
    elif sys.platform == "darwin":
        subprocess.Popen(["open", "-R", ruta] if carpeta else ["open", ruta])
    else:
        subprocess.Popen(["xdg-open", os.path.dirname(ruta) if carpeta else ruta])

def _registrar_error(detalle):
    try:
        carpeta = os.path.dirname(_ruta_cache())
        with open(os.path.join(carpeta, "error.log"), "a", encoding="utf-8") as archivo:
            archivo.write("\n%s\n%s\n" % (
                datetime.now().strftime("%Y-%m-%d %H:%M:%S"), detalle))
    except Exception:
        pass   # el registro nunca debe romper la aplicacion

class NominaApp:
    """Ventana principal con un flujo guiado de tres pasos."""

    def __init__(self, root):
        self.root = root
        self.pdfs = []
        self.instituciones = []
        self.paso = 0
        self.processing = False
        self._closing = False
        self._cola = queue.Queue()
        self._guia = None
        self.resultado = None
        self.var_destino = tk.StringVar(
            value=os.path.join(_carpeta_sugerida(), _nombre_sugerido()))
        self.var_funcionarios = tk.BooleanVar(value=INCLUIR_FUNCIONARIOS)
        self.var_abrir = tk.BooleanVar(value=True)
        self.status_var = tk.StringVar(value="Listo para comenzar.")

        self._configurar_ventana()
        self._configurar_estilos()
        self._crear_menu()
        self._crear_cabecera()
        self._crear_pasos()
        self._crear_paginas()
        self._crear_pie()
        self._ir_paso(0, silencioso=True)
        self.root.protocol("WM_DELETE_WINDOW", self._cerrar)
        self.root.bind("<Control-o>", lambda e: self._seleccionar_pdfs())
        self.root.bind("<F1>", lambda e: self._abrir_guia())
        self.root.after(800, self._chequear_actualizacion)

    # ------------------------------------------------------------------
    # Construccion de la ventana
    # ------------------------------------------------------------------
    def _configurar_ventana(self):
        self.root.title(NOMBRE_APP)
        self.root.configure(bg=COLOR_FONDO)
        self.root.minsize(880, 660)
        ancho = min(980, max(880, self.root.winfo_screenwidth() - 120))
        alto = min(760, max(660, self.root.winfo_screenheight() - 120))
        _centrar(self.root, ancho, alto)
        icono = _ruta_icono()
        if os.path.exists(icono):
            try:
                self.root.iconbitmap(icono)
            except Exception:
                pass
        self.root.option_add("*tearOff", False)

    def _configurar_estilos(self):
        estilo = ttk.Style(self.root)
        try:
            estilo.theme_use("clam")
        except tk.TclError:
            pass
        estilo.configure("TButton", font=("Segoe UI", 10), padding=(12, 6),
                         background="#E9EEF4", foreground=COLOR_TEXTO,
                         bordercolor=COLOR_BORDE, lightcolor="#E9EEF4",
                         darkcolor="#E9EEF4", relief="flat")
        estilo.map("TButton",
                   background=[("active", "#DCE6F1"), ("disabled", "#F0F3F7")],
                   foreground=[("disabled", "#9AA7B5")])
        estilo.configure("Primary.TButton", font=("Segoe UI", 10, "bold"),
                         padding=(16, 9), background=COLOR_AZUL,
                         foreground="white", bordercolor=COLOR_AZUL,
                         lightcolor=COLOR_AZUL, darkcolor=COLOR_AZUL)
        estilo.map("Primary.TButton",
                   background=[("active", COLOR_AZUL_OSCURO), ("disabled", "#AFC7DD")],
                   foreground=[("disabled", "#F4F8FC")])
        estilo.configure("Success.TButton", font=("Segoe UI", 10, "bold"),
                         padding=(16, 9), background=COLOR_VERDE,
                         foreground="white", bordercolor=COLOR_VERDE,
                         lightcolor=COLOR_VERDE, darkcolor=COLOR_VERDE)
        estilo.map("Success.TButton",
                   background=[("active", "#166336"), ("disabled", "#A8CDB8")],
                   foreground=[("disabled", "#F4F8FC")])
        estilo.configure("Treeview", font=("Segoe UI", 10), rowheight=26,
                         background="white", fieldbackground="white",
                         bordercolor=COLOR_BORDE, borderwidth=0)
        estilo.configure("Treeview.Heading", font=("Segoe UI", 9, "bold"),
                         background="#EAF1F8", foreground=COLOR_TEXTO,
                         padding=(6, 5), relief="flat")
        estilo.map("Treeview", background=[("selected", "#D6E8F7")],
                   foreground=[("selected", COLOR_TEXTO)])
        estilo.configure("TCheckbutton", font=("Segoe UI", 10),
                         background="white", foreground=COLOR_TEXTO)
        estilo.map("TCheckbutton", background=[("active", "white")])
        estilo.configure("TEntry", font=("Segoe UI", 10), padding=6)
        estilo.map("TEntry",
                   fieldbackground=[("readonly", "#F7FAFD"), ("!disabled", "white")],
                   foreground=[("readonly", COLOR_TEXTO)])
        estilo.configure("Blue.Horizontal.TProgressbar",
                         troughcolor="#E3EAF2", background=COLOR_AZUL,
                         bordercolor=COLOR_BORDE, lightcolor=COLOR_AZUL,
                         darkcolor=COLOR_AZUL, thickness=10)

    def _crear_menu(self):
        barra = tk.Menu(self.root)
        archivo = tk.Menu(barra, tearoff=0)
        archivo.add_command(label="Seleccionar PDFs…", accelerator="Ctrl+O",
                            command=self._seleccionar_pdfs)
        archivo.add_separator()
        archivo.add_command(label="Salir", command=self._cerrar)
        barra.add_cascade(label="Archivo", menu=archivo)
        ayuda = tk.Menu(barra, tearoff=0)
        ayuda.add_command(label="Guía paso a paso", accelerator="F1",
                          command=self._abrir_guia)
        ayuda.add_command(label="Acerca de", command=self._abrir_acerca)
        barra.add_cascade(label="Ayuda", menu=ayuda)
        self.root.config(menu=barra)

    def _crear_cabecera(self):
        cabecera = tk.Frame(self.root, bg=COLOR_AZUL, height=92)
        cabecera.pack(fill="x")
        cabecera.pack_propagate(False)
        izquierda = tk.Frame(cabecera, bg=COLOR_AZUL)
        izquierda.pack(side="left", padx=28, pady=14)
        tk.Label(izquierda, text=NOMBRE_APP.upper(), bg=COLOR_AZUL, fg="white",
                 font=("Segoe UI", 20, "bold")).pack(anchor="w")
        self.lbl_institucion = tk.Label(
            izquierda, text="Generador de nómina desde el Cuadro de Personal",
            bg=COLOR_AZUL, fg="#DCE9F7", font=("Segoe UI", 9),
            justify="left", anchor="w")
        self.lbl_institucion.pack(anchor="w", pady=(2, 0))

        derecha = tk.Frame(cabecera, bg=COLOR_AZUL)
        derecha.pack(side="right", padx=24)
        tk.Button(derecha, text="?  Guía de uso", command=self._abrir_guia,
                  bg=COLOR_AZUL, fg="white", activebackground=COLOR_AZUL_OSCURO,
                  activeforeground="white", bd=0, relief="flat",
                  font=("Segoe UI", 10, "bold"), cursor="hand2",
                  padx=10, pady=5).pack(side="right")
        tk.Label(derecha, text="v" + VERSION, bg=COLOR_AZUL, fg="#DCE9F7",
                 font=("Segoe UI", 9)).pack(side="right", padx=(0, 14))

    def _crear_pasos(self):
        contenedor = tk.Frame(self.root, bg="white",
                              highlightbackground=COLOR_BORDE, highlightthickness=1)
        contenedor.pack(fill="x", padx=28, pady=(14, 8))
        fila = tk.Frame(contenedor, bg="white")
        fila.pack(pady=9)
        self.step_badges = []
        self.step_labels = []
        for indice, texto in enumerate(("Archivos", "Destino", "Generar")):
            if indice:
                tk.Label(fila, text="──", bg="white", fg="#C7D2DE",
                         font=("Segoe UI", 9)).pack(side="left", padx=8)
            badge = tk.Label(fila, text=str(indice + 1), width=2, bg=COLOR_GRIS,
                             fg="white", font=("Segoe UI", 9, "bold"),
                             padx=3, pady=2)
            badge.pack(side="left")
            etiqueta = tk.Label(fila, text=texto, bg="white",
                                fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 10))
            etiqueta.pack(side="left", padx=(6, 0))
            self.step_badges.append(badge)
            self.step_labels.append(etiqueta)

    def _crear_paginas(self):
        contenedor = tk.Frame(self.root, bg=COLOR_FONDO)
        contenedor.pack(fill="both", expand=True, padx=28, pady=(0, 4))
        contenedor.grid_rowconfigure(0, weight=1)
        contenedor.grid_columnconfigure(0, weight=1)
        self.pages = {}
        constructores = (self._pagina_archivos, self._pagina_destino,
                         self._pagina_generar)
        for indice, constructor in enumerate(constructores):
            pagina = tk.Frame(contenedor, bg=COLOR_FONDO)
            pagina.grid(row=0, column=0, sticky="nsew")
            self.pages[indice] = pagina
            constructor(pagina)

    def _crear_pie(self):
        pie = tk.Frame(self.root, bg=COLOR_FONDO)
        pie.pack(fill="x", padx=28, pady=(4, 10))
        tk.Button(pie, text="¿Necesitás ayuda?  Ver guía paso a paso",
                  command=self._abrir_guia, bg=COLOR_FONDO, fg=COLOR_AZUL,
                  activebackground=COLOR_FONDO, activeforeground=COLOR_AZUL_OSCURO,
                  bd=0, relief="flat", font=("Segoe UI", 9, "underline"),
                  cursor="hand2", padx=0, pady=0).pack(side="left")
        ttk.Button(pie, text="Salir", command=self._cerrar).pack(side="right")
        tk.Label(pie, text="v" + VERSION, bg=COLOR_FONDO,
                 fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9)).pack(
                     side="right", padx=(0, 12))
        tk.Label(pie, textvariable=self.status_var, bg=COLOR_FONDO,
                 fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9), anchor="w").pack(
                     side="left", fill="x", expand=True, padx=(14, 14))

    @staticmethod
    def _tarjeta(padre):
        return tk.Frame(padre, bg="white", highlightbackground=COLOR_BORDE,
                        highlightthickness=1, bd=0)

    # ------------------------------------------------------------------
    # Pagina 1: archivos
    # ------------------------------------------------------------------
    def _pagina_archivos(self, pagina):
        tk.Label(pagina, text="Seleccioná los cuadros de personal",
                 bg=COLOR_FONDO, fg=COLOR_TEXTO,
                 font=("Segoe UI", 16, "bold")).pack(anchor="w")
        tk.Label(
            pagina,
            text="Generá la nómina de docentes y catedráticos en Excel a partir de los Cuadros "
                 "de Personal. Elegí uno o varios PDF; podés combinar el de la Escuela y el del "
                 "Colegio: las personas que aparezcan en ambos se unifican por cédula.",
            bg=COLOR_FONDO, fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 10),
            wraplength=800, justify="left",
        ).pack(anchor="w", pady=(4, 12))

        tarjeta = self._tarjeta(pagina)
        tarjeta.pack(fill="both", expand=True)

        barra = tk.Frame(tarjeta, bg="white")
        barra.pack(fill="x", padx=16, pady=(14, 8))
        self.btn_seleccionar = ttk.Button(
            barra, text="Seleccionar PDFs…", style="Primary.TButton",
            command=self._seleccionar_pdfs)
        self.btn_seleccionar.pack(side="left")
        self.btn_quitar = ttk.Button(
            barra, text="Quitar seleccionados", state="disabled",
            command=self._quitar_seleccionados)
        self.btn_quitar.pack(side="left", padx=(8, 0))
        self.btn_limpiar = ttk.Button(
            barra, text="Limpiar lista", state="disabled",
            command=self._limpiar_lista)
        self.btn_limpiar.pack(side="left", padx=(8, 0))
        self.lbl_conteo = tk.Label(barra, text="0 archivos", bg="white",
                                   fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9))
        self.lbl_conteo.pack(side="right")

        self.lista_wrap = tk.Frame(tarjeta, bg="white")
        self.lista_wrap.pack(fill="both", expand=True, padx=16, pady=(0, 6))

        self.tree_frame = tk.Frame(self.lista_wrap, bg="white")
        self.tree = ttk.Treeview(self.tree_frame, columns=("archivo", "carpeta"),
                                 show="headings", height=5, selectmode="extended")
        self.tree.heading("archivo", text="Archivo PDF")
        self.tree.heading("carpeta", text="Carpeta")
        self.tree.column("archivo", width=300, minwidth=180, stretch=False)
        self.tree.column("carpeta", width=560, minwidth=220, stretch=True)
        barra_vertical = ttk.Scrollbar(self.tree_frame, orient="vertical",
                                       command=self.tree.yview)
        self.tree.configure(yscrollcommand=barra_vertical.set)
        self.tree.pack(side="left", fill="both", expand=True)
        barra_vertical.pack(side="right", fill="y")
        self.tree.bind("<<TreeviewSelect>>", lambda e: self._actualizar_botones())
        self.tree.bind("<Delete>", lambda e: self._quitar_seleccionados())

        self.empty_frame = tk.Frame(self.lista_wrap, bg="white")
        caja = tk.Frame(self.empty_frame, bg="white")
        caja.pack(expand=True)
        tk.Label(caja, text="PDF", bg="#EAF1F8", fg=COLOR_AZUL,
                 font=("Segoe UI", 16, "bold"), width=5, height=2).pack(pady=(4, 8))
        tk.Label(caja, text="Todavía no seleccionaste ningún PDF.",
                 bg="white", fg=COLOR_TEXTO,
                 font=("Segoe UI", 11, "bold")).pack()
        tk.Label(caja, text="Hacé clic en «Seleccionar PDFs…» para comenzar.",
                 bg="white", fg=COLOR_TEXTO_SUAVE,
                 font=("Segoe UI", 9)).pack(pady=(2, 0))

        tk.Label(
            tarjeta,
            text="Consejo: los PDF deben ser los cuadros oficiales con tablas de texto. "
                 "Las fotos o escaneos no se pueden leer.",
            bg="white", fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9),
            wraplength=800, justify="left",
        ).pack(anchor="w", padx=16, pady=(0, 12))

        self.btn_siguiente_1 = ttk.Button(
            pagina, text="Siguiente: elegir destino  →", style="Primary.TButton",
            state="disabled", command=lambda: self._ir_paso(1))
        self.btn_siguiente_1.pack(side="right", pady=(10, 0))

        self.empty_frame.pack(fill="both", expand=True)
        self.tree_frame.pack_forget()

    # ------------------------------------------------------------------
    # Pagina 2: destino y opciones
    # ------------------------------------------------------------------
    def _pagina_destino(self, pagina):
        tk.Label(pagina, text="Elegí dónde guardar la nómina",
                 bg=COLOR_FONDO, fg=COLOR_TEXTO,
                 font=("Segoe UI", 16, "bold")).pack(anchor="w")
        tk.Label(
            pagina,
            text="La aplicación propone una carpeta y un nombre. Podés cambiarlos antes de generar el Excel.",
            bg=COLOR_FONDO, fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 10),
            wraplength=800, justify="left",
        ).pack(anchor="w", pady=(4, 12))

        tarjeta = self._tarjeta(pagina)
        tarjeta.pack(fill="x")
        tk.Label(tarjeta, text="Carpeta y nombre del archivo", bg="white",
                 fg=COLOR_TEXTO, font=("Segoe UI", 10, "bold")).pack(
                     anchor="w", padx=16, pady=(14, 6))
        fila = tk.Frame(tarjeta, bg="white")
        fila.pack(fill="x", padx=16, pady=(0, 6))
        self.entry_destino = ttk.Entry(fila, textvariable=self.var_destino,
                                       state="readonly")
        self.entry_destino.pack(side="left", fill="x", expand=True, ipady=3)
        self.btn_destino = ttk.Button(fila, text="Cambiar ubicación y nombre…",
                                      command=self._elegir_destino)
        self.btn_destino.pack(side="left", padx=(8, 0))
        tk.Label(tarjeta,
                 text="Si el archivo ya existe, se pedirá confirmación antes de reemplazarlo.",
                 bg="white", fg=COLOR_TEXTO_SUAVE,
                 font=("Segoe UI", 9)).pack(anchor="w", padx=16, pady=(0, 14))

        opciones = self._tarjeta(pagina)
        opciones.pack(fill="x", pady=(12, 0))
        tk.Label(opciones, text="Opciones", bg="white", fg=COLOR_TEXTO,
                 font=("Segoe UI", 10, "bold")).pack(
                     anchor="w", padx=16, pady=(12, 4))
        self.chk_funcionarios = tk.Checkbutton(
            opciones, text="Incluir funcionarios no docentes (rubros G*)",
            variable=self.var_funcionarios, bg="white", activebackground="white",
            fg=COLOR_TEXTO, activeforeground=COLOR_TEXTO, selectcolor="white",
            font=("Segoe UI", 10), bd=0, highlightthickness=0, anchor="w",
            padx=0, pady=0, cursor="hand2")
        self.chk_funcionarios.pack(anchor="w", padx=16)
        tk.Label(opciones,
                 text="Por defecto solo se incluyen docentes (L*) y catedráticos (ZZ*).",
                 bg="white", fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9)).pack(
                     anchor="w", padx=38, pady=(0, 8))
        self.chk_abrir = tk.Checkbutton(
            opciones, text="Abrir el Excel automáticamente al finalizar",
            variable=self.var_abrir, bg="white", activebackground="white",
            fg=COLOR_TEXTO, activeforeground=COLOR_TEXTO, selectcolor="white",
            font=("Segoe UI", 10), bd=0, highlightthickness=0, anchor="w",
            padx=0, pady=0, cursor="hand2")
        self.chk_abrir.pack(anchor="w", padx=16, pady=(0, 12))

        tk.Frame(pagina, bg=COLOR_FONDO).pack(fill="both", expand=True)
        navegacion = tk.Frame(pagina, bg=COLOR_FONDO)
        navegacion.pack(fill="x", pady=(12, 0))
        self.btn_atras_2 = ttk.Button(navegacion, text="←  Atrás",
                                      command=lambda: self._ir_paso(0))
        self.btn_atras_2.pack(side="left")
        self.btn_siguiente_2 = ttk.Button(
            navegacion, text="Siguiente: revisar  →", style="Primary.TButton",
            command=lambda: self._ir_paso(2))
        self.btn_siguiente_2.pack(side="right")

    # ------------------------------------------------------------------
    # Pagina 3: resumen, generacion y resultado
    # ------------------------------------------------------------------
    def _pagina_generar(self, pagina):
        tk.Label(pagina, text="Revisá y generá la nómina",
                 bg=COLOR_FONDO, fg=COLOR_TEXTO,
                 font=("Segoe UI", 16, "bold")).pack(anchor="w")
        tk.Label(
            pagina,
            text="Verificá los archivos y el destino. Después, hacé clic en «Generar nómina ahora».",
            bg=COLOR_FONDO, fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 10),
            wraplength=800, justify="left",
        ).pack(anchor="w", pady=(4, 12))

        resumen = self._tarjeta(pagina)
        resumen.pack(fill="x")
        resumen.grid_columnconfigure(1, weight=1)
        tk.Label(resumen, text="Resumen", bg="white", fg=COLOR_TEXTO,
                 font=("Segoe UI", 10, "bold")).grid(
                     row=0, column=0, columnspan=2, sticky="w", padx=16, pady=(12, 6))
        tk.Label(resumen, text="Archivos", bg="white", fg=COLOR_TEXTO_SUAVE,
                 font=("Segoe UI", 9, "bold")).grid(
                     row=1, column=0, sticky="nw", padx=(16, 10))
        self.lbl_resumen_archivos = tk.Label(
            resumen, text="", bg="white", fg=COLOR_TEXTO, font=("Segoe UI", 9),
            justify="left", anchor="w", wraplength=740)
        self.lbl_resumen_archivos.grid(row=1, column=1, sticky="w",
                                       padx=(0, 16), pady=(0, 4))
        tk.Label(resumen, text="Destino", bg="white", fg=COLOR_TEXTO_SUAVE,
                 font=("Segoe UI", 9, "bold")).grid(
                     row=2, column=0, sticky="nw", padx=(16, 10))
        self.lbl_resumen_destino = tk.Label(
            resumen, text="", bg="white", fg=COLOR_TEXTO, font=("Segoe UI", 9),
            justify="left", anchor="w", wraplength=740)
        self.lbl_resumen_destino.grid(row=2, column=1, sticky="w",
                                      padx=(0, 16), pady=(0, 4))
        tk.Label(resumen, text="Opciones", bg="white", fg=COLOR_TEXTO_SUAVE,
                 font=("Segoe UI", 9, "bold")).grid(
                     row=3, column=0, sticky="nw", padx=(16, 10))
        self.lbl_resumen_opciones = tk.Label(
            resumen, text="", bg="white", fg=COLOR_TEXTO, font=("Segoe UI", 9),
            justify="left", anchor="w")
        self.lbl_resumen_opciones.grid(row=3, column=1, sticky="w",
                                       padx=(0, 16), pady=(0, 12))

        self.accion_wrap = tk.Frame(pagina, bg=COLOR_FONDO)
        self.accion_wrap.pack(fill="x", pady=(12, 0))

        self.action_card = self._tarjeta(self.accion_wrap)
        self.action_card.pack(fill="x")
        accion = tk.Frame(self.action_card, bg="white")
        accion.pack(fill="x", padx=16, pady=14)
        self.btn_generar = ttk.Button(
            accion, text="Generar nómina ahora", style="Primary.TButton",
            command=self._generar)
        self.btn_generar.pack(fill="x")
        self.progress = ttk.Progressbar(
            accion, mode="determinate", maximum=100, value=0,
            style="Blue.Horizontal.TProgressbar")
        self.progress.pack(fill="x", pady=(10, 6))
        self.lbl_estado = tk.Label(
            accion, text="Listo para generar.", bg="white",
            fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9), anchor="w")
        self.lbl_estado.pack(fill="x")

        self.result_card = tk.Frame(
            self.accion_wrap, bg=COLOR_VERDE_CLARO,
            highlightbackground=COLOR_VERDE, highlightthickness=1)
        contenido = tk.Frame(self.result_card, bg=COLOR_VERDE_CLARO)
        contenido.pack(fill="x", padx=16, pady=14)
        tk.Label(contenido, text="✓  Nómina generada correctamente",
                 bg=COLOR_VERDE_CLARO, fg=COLOR_VERDE,
                 font=("Segoe UI", 12, "bold")).pack(anchor="w")
        self.stats_frame = tk.Frame(contenido, bg=COLOR_VERDE_CLARO)
        self.stats_frame.pack(fill="x", pady=(10, 4))
        self.lbl_resultado = tk.Label(
            contenido, text="", bg=COLOR_VERDE_CLARO, fg=COLOR_TEXTO_SUAVE,
            font=("Segoe UI", 9), justify="left", anchor="w", wraplength=780)
        self.lbl_resultado.pack(fill="x", pady=(0, 8))
        botones = tk.Frame(contenido, bg=COLOR_VERDE_CLARO)
        botones.pack(fill="x")
        ttk.Button(botones, text="Abrir Excel", style="Success.TButton",
                   command=self._abrir_excel).pack(side="left")
        ttk.Button(botones, text="Abrir carpeta",
                   command=self._abrir_carpeta).pack(side="left", padx=(8, 0))
        ttk.Button(botones, text="Generar de nuevo",
                   command=self._mostrar_accion).pack(side="left", padx=(8, 0))

        tk.Frame(pagina, bg=COLOR_FONDO).pack(fill="both", expand=True)
        navegacion = tk.Frame(pagina, bg=COLOR_FONDO)
        navegacion.pack(fill="x", pady=(10, 0))
        self.btn_atras_3 = ttk.Button(navegacion, text="←  Atrás",
                                      command=lambda: self._ir_paso(1))
        self.btn_atras_3.pack(side="left")

    # ------------------------------------------------------------------
    # Navegacion y estado
    # ------------------------------------------------------------------
    def _ir_paso(self, paso, silencioso=False):
        if paso > 0 and not self.pdfs:
            if not silencioso:
                messagebox.showwarning(
                    "Faltan archivos",
                    "Seleccioná al menos un PDF del cuadro de personal.",
                    parent=self.root)
            return
        if paso == 2:
            if not self.var_destino.get().strip():
                self._elegir_destino()
                if not self.var_destino.get().strip():
                    return
            self._actualizar_resumen()
        self.paso = paso
        self.pages[paso].tkraise()
        self._actualizar_paso(paso)
        self._actualizar_botones()
        if not silencioso:
            if paso == 0:
                self.status_var.set("Paso 1 de 3: seleccioná los PDF del cuadro de personal.")
            elif paso == 1:
                self.status_var.set("Paso 2 de 3: elegí dónde guardar la nómina.")
            else:
                self.status_var.set("Paso 3 de 3: revisá el resumen y generá la nómina.")

    def _actualizar_paso(self, actual):
        for indice, (badge, etiqueta) in enumerate(
                zip(self.step_badges, self.step_labels)):
            if indice < actual:
                badge.config(text="✓", bg=COLOR_VERDE)
                etiqueta.config(fg=COLOR_VERDE, font=("Segoe UI", 10, "bold"))
            elif indice == actual:
                badge.config(text=str(indice + 1), bg=COLOR_AZUL)
                etiqueta.config(fg=COLOR_AZUL, font=("Segoe UI", 10, "bold"))
            else:
                badge.config(text=str(indice + 1), bg=COLOR_GRIS)
                etiqueta.config(fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 10))

    def _actualizar_botones(self):
        if self.processing:
            return
        self.btn_siguiente_1.state(["!disabled"] if self.pdfs else ["disabled"])
        self.btn_siguiente_2.state(
            ["!disabled"] if self.var_destino.get().strip() else ["disabled"])
        self.btn_quitar.state(
            ["!disabled"] if self.tree.selection() else ["disabled"])
        self.btn_limpiar.state(["!disabled"] if self.pdfs else ["disabled"])
        self.btn_generar.state(
            ["!disabled"] if self.pdfs and self.var_destino.get().strip()
            else ["disabled"])

    def _actualizar_lista(self):
        if self.pdfs:
            self.empty_frame.pack_forget()
            self.tree_frame.pack(fill="both", expand=True)
        else:
            self.tree_frame.pack_forget()
            self.empty_frame.pack(fill="both", expand=True)
        self.lbl_conteo.config(text="%d archivo%s" % (
            len(self.pdfs), "" if len(self.pdfs) == 1 else "s"))
        self._actualizar_instituciones()
        self._actualizar_botones()

    def _mostrar_instituciones(self):
        nombres = [info["nombre"] for info in self.instituciones if info.get("nombre")]
        if nombres:
            self.lbl_institucion.config(text="  ·  ".join(nombres))
            if len(nombres) == 1:
                self.root.title("%s — %s" % (NOMBRE_APP, nombres[0]))
            else:
                self.root.title("%s — %d instituciones" % (NOMBRE_APP, len(nombres)))
        else:
            self.lbl_institucion.config(
                text="Generador de nómina desde el Cuadro de Personal")
            self.root.title(NOMBRE_APP)

    def _actualizar_instituciones(self):
        if self.pdfs:
            try:
                self.instituciones = extraer_instituciones(self.pdfs)
            except Exception:
                self.instituciones = []
        else:
            self.instituciones = []
        self._mostrar_instituciones()

    def _actualizar_resumen(self):
        if self.pdfs:
            nombres = [os.path.basename(p) for p in self.pdfs]
            self.lbl_resumen_archivos.config(text="%d archivo%s: %s" % (
                len(nombres), "" if len(nombres) == 1 else "s",
                ", ".join(nombres)))
        else:
            self.lbl_resumen_archivos.config(text="Ningún archivo seleccionado.")
        self.lbl_resumen_destino.config(
            text=self.var_destino.get() or "(sin destino)")
        self.lbl_resumen_opciones.config(
            text="Incluir funcionarios (G*): %s   ·   Abrir al finalizar: %s" % (
                "Sí" if self.var_funcionarios.get() else "No",
                "Sí" if self.var_abrir.get() else "No"))

    def _set_busy(self, ocupado):
        self.processing = ocupado
        widgets = (
            self.btn_seleccionar, self.btn_quitar, self.btn_limpiar,
            self.btn_siguiente_1, self.btn_siguiente_2, self.btn_generar,
            self.btn_destino, self.btn_atras_2, self.btn_atras_3,
            self.chk_funcionarios, self.chk_abrir,
        )
        for widget in widgets:
            if isinstance(widget, ttk.Widget):
                widget.state(["disabled"] if ocupado else ["!disabled"])
            else:
                widget.configure(state="disabled" if ocupado else "normal")
        if not ocupado:
            self._actualizar_botones()

    # ------------------------------------------------------------------
    # Seleccion de archivos y destino
    # ------------------------------------------------------------------
    def _seleccionar_pdfs(self):
        if self.processing:
            return
        rutas = filedialog.askopenfilenames(
            parent=self.root,
            title="Seleccioná el/los cuadro(s) de personal (PDF)",
            filetypes=[("Archivos PDF", "*.pdf")])
        if rutas:
            self._agregar_pdfs(rutas)

    def _agregar_pdfs(self, rutas):
        agregados = 0
        for ruta in rutas:
            ruta = os.path.abspath(ruta)
            if not ruta.lower().endswith(".pdf") or not os.path.isfile(ruta):
                continue
            if ruta in self.pdfs:
                continue
            self.pdfs.append(ruta)
            self.tree.insert("", "end", values=(
                os.path.basename(ruta), os.path.dirname(ruta)))
            agregados += 1
        if agregados:
            self.resultado = None
            self._mostrar_accion()
            self._actualizar_lista()
            self.status_var.set("%d archivo%s listo%s para procesar." % (
                len(self.pdfs), "" if len(self.pdfs) == 1 else "s",
                "" if len(self.pdfs) == 1 else "s"))
        else:
            self.status_var.set(
                "Los archivos elegidos ya estaban en la lista o no son PDF válidos.")
        return agregados

    def _quitar_seleccionados(self):
        if self.processing:
            return
        seleccion = self.tree.selection()
        if not seleccion:
            return
        items = list(self.tree.get_children())
        indices = sorted((items.index(item) for item in seleccion), reverse=True)
        for indice in indices:
            self.tree.delete(items[indice])
            del self.pdfs[indice]
        self.resultado = None
        self._mostrar_accion()
        self._actualizar_lista()
        self.status_var.set("Archivo(s) quitado(s) de la lista.")

    def _limpiar_lista(self):
        if self.processing or not self.pdfs:
            return
        if not messagebox.askyesno(
                "Limpiar lista",
                "¿Querés quitar los %d archivos de la lista?" % len(self.pdfs),
                parent=self.root):
            return
        self.pdfs = []
        for item in self.tree.get_children():
            self.tree.delete(item)
        self.resultado = None
        self._mostrar_accion()
        self._actualizar_lista()
        self.status_var.set("Lista vacía. Seleccioná los PDF para comenzar.")

    def _elegir_destino(self):
        if self.processing:
            return
        actual = self.var_destino.get().strip()
        carpeta = os.path.dirname(actual) if actual else _carpeta_sugerida()
        if not os.path.isdir(carpeta):
            carpeta = _carpeta_sugerida()
        nombre = os.path.basename(actual) if actual else _nombre_sugerido()
        ruta = filedialog.asksaveasfilename(
            parent=self.root, title="Guardar la nómina como...",
            defaultextension=".xlsx", initialdir=carpeta, initialfile=nombre,
            filetypes=[("Libro de Excel", "*.xlsx")])
        if not ruta:
            return
        if not ruta.lower().endswith(".xlsx"):
            ruta += ".xlsx"
        self.var_destino.set(ruta)
        self.resultado = None
        self._mostrar_accion()
        self._actualizar_resumen()
        self._actualizar_botones()
        self.status_var.set("Destino actualizado.")

    # ------------------------------------------------------------------
    # Generacion en segundo plano
    # ------------------------------------------------------------------
    def _generar(self):
        if self.processing:
            return
        if not self.pdfs:
            messagebox.showwarning(
                "Faltan archivos",
                "Seleccioná al menos un PDF del cuadro de personal.",
                parent=self.root)
            self._ir_paso(0)
            return
        destino = self.var_destino.get().strip()
        if not destino:
            self._elegir_destino()
            destino = self.var_destino.get().strip()
            if not destino:
                return
        if not destino.lower().endswith(".xlsx"):
            destino += ".xlsx"
            self.var_destino.set(destino)
        carpeta = os.path.dirname(destino)
        if not carpeta or not os.path.isdir(carpeta):
            messagebox.showerror(
                "Carpeta no encontrada",
                "La carpeta de destino no existe:\n\n%s\n\nElegí otra ubicación." % carpeta,
                parent=self.root)
            return
        if os.path.exists(destino):
            if not messagebox.askyesno(
                    "Reemplazar archivo",
                    "Ya existe un archivo con ese nombre:\n\n%s\n\n¿Querés reemplazarlo?" % destino,
                    parent=self.root):
                return
        self._mostrar_accion()
        self._set_busy(True)
        self.resultado = None
        self.progress.configure(mode="indeterminate", value=0)
        self.progress.start(12)
        self.lbl_estado.config(text="Preparando la nómina...", fg=COLOR_TEXTO_SUAVE)
        self.status_var.set("Procesando los PDF. No cierres la ventana...")
        hilo = threading.Thread(
            target=self._trabajar,
            args=(list(self.pdfs), destino, self.var_funcionarios.get()),
            daemon=True)
        hilo.start()
        self.root.after(80, self._revisar_cola)

    def _trabajar(self, pdfs, destino, incluir_funcionarios):
        try:
            resumen = procesar(
                pdfs, destino, incluir_funcionarios=incluir_funcionarios,
                progreso=self._informar_progreso)
            self._cola.put(("ok", resumen))
        except Exception as error:
            self._cola.put(("error", error, traceback.format_exc()))

    def _informar_progreso(self, actual, total, mensaje):
        self._cola.put(("progreso", actual, total, mensaje))

    def _revisar_cola(self):
        if self._closing:
            return
        try:
            if not self.root.winfo_exists():
                return
        except tk.TclError:
            return
        try:
            while True:
                mensaje = self._cola.get_nowait()
                tipo = mensaje[0]
                if tipo == "progreso":
                    _, actual, total, texto = mensaje
                    self.lbl_estado.config(text=texto)
                    if total:
                        self.status_var.set("%s  (%d/%d)" % (texto, actual, total))
                    else:
                        self.status_var.set(texto)
                elif tipo == "ok":
                    self._terminar_ok(mensaje[1])
                    return
                elif tipo == "error":
                    self._terminar_error(mensaje[1], mensaje[2])
                    return
        except queue.Empty:
            pass
        self.root.after(80, self._revisar_cola)

    def _terminar_ok(self, resumen):
        self.progress.stop()
        self.progress.configure(mode="determinate", value=100)
        self._set_busy(False)
        self.resultado = resumen
        if "instituciones" in resumen:
            self.instituciones = resumen["instituciones"]
            self._mostrar_instituciones()
        self.lbl_estado.config(text="Nómina generada correctamente.",
                               fg=COLOR_VERDE)
        self.status_var.set("Listo. Nómina guardada correctamente.")
        self._mostrar_resultado(resumen)
        if self.var_abrir.get():
            self.root.after(250, lambda: self._abrir_excel(mostrar_error=False))

    def _terminar_error(self, error, detalle):
        self.progress.stop()
        self.progress.configure(mode="determinate", value=0)
        self._set_busy(False)
        self.lbl_estado.config(text="No se pudo generar la nómina.",
                               fg="#B42318")
        self.status_var.set("Error al generar la nómina.")
        _registrar_error(detalle)
        if isinstance(error, PermissionError):
            texto = (
                "No se pudo guardar el archivo.\n\n"
                "Si el Excel está abierto, cerralo y volvé a intentar.\n\n"
                "También verificá que tengas permiso para escribir en esa carpeta.\n\n"
                "Detalle: %s" % error)
        else:
            texto = (
                "Ocurrió un error al generar la nómina.\n\n"
                "Revisá que los PDF sean los cuadros de personal oficiales y que no estén dañados.\n\n"
                "Detalle: %s" % error)
        messagebox.showerror("No se pudo generar la nómina", texto,
                             parent=self.root)

    def _poner_estadisticas(self, resumen):
        for widget in self.stats_frame.winfo_children():
            widget.destroy()
        datos = [("Docentes", resumen["docentes"]),
                 ("Catedráticos", resumen["catedraticos"])]
        if resumen.get("funcionarios"):
            datos.append(("Funcionarios", resumen["funcionarios"]))
        if resumen.get("otros"):
            datos.append(("Sin categoría", len(resumen["otros"])))
        datos.append(("Total", resumen["total"]))
        for etiqueta, valor in datos:
            bloque = tk.Frame(self.stats_frame, bg=COLOR_VERDE_CLARO)
            bloque.pack(side="left", padx=(0, 28))
            tk.Label(bloque, text=str(valor), bg=COLOR_VERDE_CLARO,
                     fg=COLOR_VERDE, font=("Segoe UI", 20, "bold")).pack(anchor="w")
            tk.Label(bloque, text=etiqueta, bg=COLOR_VERDE_CLARO,
                     fg=COLOR_TEXTO_SUAVE, font=("Segoe UI", 9)).pack(anchor="w")

    def _mostrar_resultado(self, resumen):
        self._poner_estadisticas(resumen)
        notas = []
        excluidos = resumen.get("excluidos", 0)
        if excluidos == 1:
            notas.append("Se excluyó 1 persona que no es docente ni catedrática.")
        elif excluidos > 1:
            notas.append("Se excluyeron %d personas que no son docentes ni catedráticas." % excluidos)
        otros = len(resumen.get("otros", []))
        if otros == 1:
            notas.append("Atención: 1 persona tiene una categoría no reconocida; revisá el Excel.")
        elif otros > 1:
            notas.append("Atención: %d personas tienen categorías no reconocidas; revisá el Excel." % otros)
        texto = "Guardado en:\n%s" % resumen["destino"]
        if notas:
            texto += "\n\n" + "  ".join(notas)
        self.lbl_resultado.config(text=texto)
        self.action_card.pack_forget()
        if not self.result_card.winfo_manager():
            self.result_card.pack(fill="x")

    def _mostrar_accion(self):
        if self.processing:
            return
        self.result_card.pack_forget()
        if not self.action_card.winfo_manager():
            self.action_card.pack(fill="x")
        self.progress.configure(mode="determinate", value=0)
        self.lbl_estado.config(text="Listo para generar.", fg=COLOR_TEXTO_SUAVE)

    # ------------------------------------------------------------------
    # Abrir resultados
    # ------------------------------------------------------------------
    def _abrir_excel(self, mostrar_error=True):
        if not self.resultado:
            return
        try:
            _abrir_con_sistema(self.resultado["destino"], carpeta=False)
        except Exception as error:
            if mostrar_error:
                messagebox.showerror(
                    "No se pudo abrir el Excel",
                    "Abrí el archivo manualmente desde:\n\n%s\n\nDetalle: %s" % (
                        self.resultado["destino"], error),
                    parent=self.root)

    def _abrir_carpeta(self):
        if not self.resultado:
            return
        try:
            _abrir_con_sistema(self.resultado["destino"], carpeta=True)
        except Exception as error:
            messagebox.showerror(
                "No se pudo abrir la carpeta",
                "La nómina está en:\n\n%s\n\nDetalle: %s" % (
                    self.resultado["destino"], error),
                parent=self.root)

    # ------------------------------------------------------------------
    # Ayuda y actualizaciones
    # ------------------------------------------------------------------
    def _abrir_guia(self):
        if self._guia is not None:
            try:
                if self._guia.winfo_exists():
                    self._guia.lift()
                    self._guia.focus_force()
                    return
            except tk.TclError:
                pass
        dialogo = tk.Toplevel(self.root)
        self._guia = dialogo
        dialogo.title("Guía paso a paso — %s" % NOMBRE_APP)
        dialogo.configure(bg="white")
        dialogo.transient(self.root)
        dialogo.grab_set()
        dialogo.minsize(640, 480)
        _centrar(dialogo, 780, 640, padre=self.root)

        cabecera = tk.Frame(dialogo, bg=COLOR_AZUL, height=64)
        cabecera.pack(fill="x")
        cabecera.pack_propagate(False)
        tk.Label(cabecera, text="Guía paso a paso", bg=COLOR_AZUL, fg="white",
                 font=("Segoe UI", 15, "bold")).pack(side="left", padx=20)
        tk.Label(cabecera, text="Todo lo necesario para generar la nómina",
                 bg=COLOR_AZUL, fg="#DCE9F7",
                 font=("Segoe UI", 9)).pack(side="left", padx=(0, 20))

        cuerpo = tk.Frame(dialogo, bg="white")
        cuerpo.pack(fill="both", expand=True)
        texto = tk.Text(
            cuerpo, wrap="word", bg="white", fg=COLOR_TEXTO,
            font=("Segoe UI", 10), padx=22, pady=16, bd=0,
            highlightthickness=0, cursor="arrow")
        barra = ttk.Scrollbar(cuerpo, orient="vertical", command=texto.yview)
        texto.configure(yscrollcommand=barra.set)
        texto.pack(side="left", fill="both", expand=True)
        barra.pack(side="right", fill="y")
        texto.bind("<MouseWheel>", lambda e: texto.yview_scroll(
            int(-1 * (e.delta / 120)), "units"))

        texto.tag_configure("h1", font=("Segoe UI", 16, "bold"),
                            foreground=COLOR_AZUL, spacing3=6)
        texto.tag_configure("h2", font=("Segoe UI", 11, "bold"),
                            foreground=COLOR_AZUL, spacing1=14, spacing3=4)
        texto.tag_configure("p", spacing3=6)
        texto.tag_configure("bullet", lmargin1=18, lmargin2=32, spacing3=4)
        texto.tag_configure("nota", font=("Segoe UI", 9),
                            foreground=COLOR_TEXTO_SUAVE,
                            spacing1=8, spacing3=4)

        def escribir(etiqueta, contenido):
            texto.insert("end", contenido, etiqueta)

        escribir("h1", "¿Qué hace esta aplicación?\n")
        escribir("p", "Genera la nómina de docentes y catedráticos en Excel a partir de los "
                      "Cuadros de Personal en PDF de la Escuela y del Colegio. Las personas "
                      "repetidas se unifican por cédula y el archivo queda listo para imprimir.\n")

        escribir("h2", "Antes de empezar\n")
        escribir("bullet", "• Tené a mano los PDF del Cuadro de Personal de la Escuela y/o del Colegio.\n")
        escribir("bullet", "• Los PDF deben contener tablas de texto. Las fotos o escaneos no se pueden leer.\n")
        escribir("bullet", "• La aplicación no modifica los PDF originales.\n")

        escribir("h2", "Paso 1 — Seleccionar los PDF\n")
        escribir("bullet", "• Hacé clic en «Seleccionar PDFs…».\n")
        escribir("bullet", "• Elegí uno o varios archivos. Con Ctrl o Shift podés seleccionar más de uno.\n")
        escribir("bullet", "• Revisá la lista. Si te equivocaste, seleccioná uno o varios archivos y usá «Quitar seleccionados».\n")
        escribir("bullet", "• Cuando esté todo, hacé clic en «Siguiente: elegir destino».\n")

        escribir("h2", "Paso 2 — Elegir dónde guardar\n")
        escribir("bullet", "• La aplicación propone una carpeta y un nombre de archivo.\n")
        escribir("bullet", "• Para cambiarlos, hacé clic en «Cambiar ubicación y nombre…».\n")
        escribir("bullet", "• Opcional: marcá «Incluir funcionarios no docentes (G*)» si también los necesitás en la nómina.\n")
        escribir("bullet", "• Después, hacé clic en «Siguiente: revisar».\n")

        escribir("h2", "Paso 3 — Generar la nómina\n")
        escribir("bullet", "• Revisá el resumen de archivos, destino y opciones.\n")
        escribir("bullet", "• Hacé clic en «Generar nómina ahora» y esperá mientras se leen los PDF.\n")
        escribir("bullet", "• Al terminar vas a ver el total de docentes, catedráticos y funcionarios.\n")
        escribir("bullet", "• Desde ahí podés abrir el Excel, abrir la carpeta o generar otra nómina.\n")

        escribir("h2", "Qué hace con los datos\n")
        escribir("bullet", "• Unifica por cédula: si una persona aparece en varios PDF o turnos, queda una sola fila.\n")
        escribir("bullet", "• Muestra todos los turnos y rubros de cada persona.\n")
        escribir("bullet", "• Ordena la lista alfabéticamente por apellido.\n")
        escribir("bullet", "• Descarta vacantes y filas sin una cédula válida de 4 a 8 dígitos.\n")

        escribir("h2", "Preguntas frecuentes\n")
        escribir("p", "¿Por qué no aparece una persona?\n")
        escribir("nota", "Verificá que la fila esté en una tabla del PDF y que la cédula tenga entre 4 y 8 dígitos. "
                         "Las filas con «VACANTE» o sin cédula se descartan.\n")
        escribir("p", "¿Por qué un nombre quedó mal?\n")
        escribir("nota", "La separación de nombres es automática. Los casos ambiguos se corrigen desde "
                         "manual_nombres.py, un archivo técnico opcional que no se distribuye con la aplicación.\n")
        escribir("p", "¿Puedo volver a generar la nómina?\n")
        escribir("nota", "Sí. Usá «Generar de nuevo» o volvé atrás para cambiar archivos, destino u opciones. "
                         "Si el archivo está abierto en Excel, cerralo antes de reemplazarlo.\n")
        escribir("p", "¿Dónde queda el archivo?\n")
        escribir("nota", "En la ruta que elegiste en el paso 2. También podés abrir la carpeta desde el resultado.\n")

        escribir("h2", "Privacidad\n")
        escribir("p", "Los PDF y el Excel contienen nombres y cédulas. Guardalos en un lugar seguro y no los subas "
                      "a internet. La aplicación no envía los PDF ni el Excel a ningún servidor; solo consulta "
                      "GitHub para buscar actualizaciones, sin enviar datos personales.\n")
        texto.config(state="disabled")

        pie = tk.Frame(dialogo, bg="#F7FAFD",
                       highlightbackground=COLOR_BORDE, highlightthickness=1)
        pie.pack(fill="x")
        ttk.Button(pie, text="Cerrar", command=dialogo.destroy).pack(
            side="right", padx=16, pady=10)

        def ir_a_seleccion():
            dialogo.destroy()
            self._ir_paso(0)
            self.root.after(50, self._seleccionar_pdfs)

        ttk.Button(pie, text="Seleccionar PDFs ahora", style="Primary.TButton",
                   command=ir_a_seleccion).pack(side="right", padx=(0, 8), pady=10)
        dialogo.bind("<Escape>", lambda e: dialogo.destroy())
        dialogo.protocol("WM_DELETE_WINDOW", dialogo.destroy)
        dialogo.focus_set()

    def _abrir_acerca(self):
        messagebox.showinfo(
            "Acerca de %s" % NOMBRE_APP,
            "%s  v%s\n\n"
            "Genera la nómina de docentes y catedráticos en Excel a partir de los "
            "Cuadros de Personal en PDF.\n\n"
            "Autor: %s\n"
            "Repositorio:\nhttps://github.com/%s" % (
                NOMBRE_APP, VERSION, AUTOR, REPO),
            parent=self.root)

    def _chequear_actualizacion(self):
        if self._closing:
            return
        try:
            if verificar_actualizacion(self.root):
                self._closing = True
                self.root.destroy()
        except Exception:
            pass

    def _cerrar(self):
        if self.processing:
            if not messagebox.askyesno(
                    "Salir",
                    "La nómina se está generando. Si salís ahora, el archivo puede quedar incompleto.\n\n"
                    "¿Querés salir igual?",
                    parent=self.root):
                return
        self._closing = True
        self.root.destroy()

def _excepcion_no_controlada(tipo, valor, traza):
    detalle = "".join(traceback.format_exception(tipo, valor, traza))
    _registrar_error(detalle)
    try:
        messagebox.showerror(
            "Error inesperado",
            "Ocurrió un error inesperado en la aplicación:\n\n%s\n\n"
            "Podés cerrar y volver a intentar." % valor)
    except Exception:
        pass

def main_gui():
    if tk is None:
        raise SystemExit(
            "No se encontró tkinter. Instalá Tk para usar la interfaz gráfica "
            "o ejecutá el modo consola con --cli.")
    root = tk.Tk()
    root.report_callback_exception = _excepcion_no_controlada
    NominaApp(root)
    root.mainloop()

# ----------------------------------------------------------------------
# 8. Modo consola
# ----------------------------------------------------------------------
def find_pdf(keyword):
    candidates = []
    for root, dirs, files in os.walk(BASE):
        depth = root[len(BASE):].count(os.sep)
        if depth > 2:
            continue
        for f in files:
            if f.lower().endswith(".pdf") and keyword.lower() in f.lower():
                candidates.append(os.path.join(root, f))
    candidates.sort(key=lambda p: os.path.getmtime(p), reverse=True)
    return candidates[0] if candidates else None

def main_cli(argv):
    pdfs, out, todos = [], None, False
    if "--pdfs" in argv:
        i = argv.index("--pdfs")
        while i + 1 < len(argv) and not argv[i + 1].startswith("--"):
            pdfs.append(argv[i + 1])
            i += 1
    if "--out" in argv:
        out = argv[argv.index("--out") + 1]
    todos = "--todos" in argv

    if not pdfs:  # fallback: PDFs de la carpeta estandar
        for kw in ("ESCUELA", "COLEGIO"):
            p = find_pdf(kw)
            if p:
                pdfs.append(p)
    if not pdfs:
        raise FileNotFoundError("No se encontraron PDFs del cuadro de personal")
    if not out:
        out = os.path.join(BASE, "NOMINA DOCENTES ESCUELA Y COLEGIO 2026 - GENERADO.xlsx")

    print("PDFs:")
    for p in pdfs:
        print("  -", p)

    resumen = procesar(pdfs, out, incluir_funcionarios=todos)

    print()
    print("Generado:", resumen["destino"])
    print("Version del script:", VERSION)
    print("Total docentes:", resumen["total"])
    print("  Docentes (L*):", resumen["docentes"])
    print("  Catedraticos (ZZ*):", resumen["catedraticos"])
    print("  Excluidos funcionarios (G*):", resumen["excluidos"])
    for x in resumen["otros"]:
        print("  !! Revisar categoria:", x["nombre"], x["ced"], x["categoria"])

if __name__ == "__main__":
    argv = sys.argv[1:]
    if "--cli" in argv or "--pdfs" in argv or "--out" in argv or "--todos" in argv:
        main_cli(argv)
    else:
        main_gui()
